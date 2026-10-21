"""Parsers for every document type in the Citi accruals process.

All parsing is deterministic (no AI): Citi invoices, the term sheets and the FA register all have
consistent text layouts, so pattern matching gives exact, auditable results.
"""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date

import pandas as pd
import pdfplumber

from .dates import parse_date


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def pdf_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join((page.extract_text() or "") for page in pdf.pages)


def _money(s: str) -> float:
    return float(s.replace(",", ""))


# --------------------------------------------------------------------------------------------
# Citi invoice
# --------------------------------------------------------------------------------------------
@dataclass
class InvoiceLine:
    line_no: int
    series: str
    maturity_date: date
    action_date: date
    amount: float
    currency: str
    description: str = "Trustee per Series per year"


@dataclass
class ParsedInvoice:
    invoice_no: str
    program_code: str
    program_text: str
    invoice_date: date | None
    arrears_start: date | None
    arrears_end: date | None
    advance_start: date | None
    advance_end: date | None
    fee_total: float | None
    currency: str
    lines: list[InvoiceLine] = field(default_factory=list)
    approval_date: date | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def lines_total(self) -> float:
        return round(sum(l.amount for l in self.lines), 2)

    @property
    def totals_tie(self) -> bool:
        return self.fee_total is not None and abs(self.lines_total - self.fee_total) < 0.005


_LINE_RE = re.compile(
    r"Series\s+(?P<series>\d{4}-\d+)\s+Maturing\s+(?P<mat>\d{1,2}-[A-Za-z]{3}-\d{2,4})\s+"
    r"(?P<act>\d{1,2}-[A-Za-z]{3}-\d{2,4})\s+(?P<amt>[\d,]+\.\d{2})\s+(?P<ccy>[A-Z]{3})"
)
_PERIOD_RE = r"(\d{1,2}[ -][A-Za-z]{3}[ -]\d{4})\s*-\s*(\d{1,2}[ -][A-Za-z]{3}[ -]\d{4})"


def _period(text: str, label: str) -> tuple[date | None, date | None]:
    m = re.search(label + r":\s*" + _PERIOD_RE, text)
    if not m:
        return None, None
    return parse_date(m.group(1).replace("-", " ")), parse_date(m.group(2).replace("-", " "))


def looks_like_citi_invoice(text: str) -> bool:
    return "INVOICE" in text and "Invoice Number" in text and "REPACK" in text


def parse_invoice(data: bytes) -> ParsedInvoice:
    text = pdf_text(data)
    return parse_invoice_text(text)


def parse_invoice_text(text: str) -> ParsedInvoice:
    warnings: list[str] = []
    inv_no = re.search(r"Invoice Number:\s*(\d+)", text)
    if not inv_no:
        raise ValueError("This PDF does not look like a Citi invoice (no 'Invoice Number' found).")
    prog = re.search(r"(REPACK\d+)\s+(.+)", text)
    inv_date = re.search(r"Invoice Date:\s*([A-Za-z]+ \d{1,2}, \d{4}|\d{1,2}-[A-Za-z]{3}-\d{4})", text)
    total = re.search(r"(?:Fee Total|Total Invoice Amount)\s+([\d,]+\.\d{2})\s*([A-Z]{3})?", text)
    arrears = _period(text, "Arrears Period")
    advance = _period(text, "Advance Period")

    lines = []
    for i, m in enumerate(_LINE_RE.finditer(text), start=1):
        lines.append(InvoiceLine(
            line_no=i,
            series=m.group("series"),
            maturity_date=parse_date(m.group("mat")),
            action_date=parse_date(m.group("act")),
            amount=_money(m.group("amt")),
            currency=m.group("ccy"),
        ))
    # Approval stamp added by the business, e.g. "okay to pay / <approver name> - 1/20/2026". Only the date is
    # kept: the approver recorded in the app is the signed-in (SSO) user who posts the approval.
    approval_date = None
    stamp = re.search(r"(?i)ok(?:ay)?\s+to\s+pay\s*\n?\s*([A-Za-z][A-Za-z .'-]+?)\s*-\s*(\d{1,2}/\d{1,2}/\d{4})", text)
    if stamp:
        approval_date = parse_date(stamp.group(2))

    if not lines:
        warnings.append("No fee lines were found on the invoice.")
    if not total:
        warnings.append("Invoice total could not be read.")
    if not prog:
        warnings.append("Citi program code (REPACK####) not found.")

    return ParsedInvoice(
        invoice_no=inv_no.group(1),
        program_code=prog.group(1) if prog else "UNKNOWN",
        program_text=prog.group(2).strip() if prog else "",
        invoice_date=parse_date(inv_date.group(1)) if inv_date else None,
        arrears_start=arrears[0], arrears_end=arrears[1],
        advance_start=advance[0], advance_end=advance[1],
        fee_total=_money(total.group(1)) if total else None,
        currency=(total.group(2) if total and total.group(2) else "USD"),
        lines=lines,
        approval_date=approval_date,
        warnings=warnings,
    )


# --------------------------------------------------------------------------------------------
# Term sheet
# --------------------------------------------------------------------------------------------
@dataclass
class ParsedTermSheet:
    series: str | None
    issuer: str | None
    trade_date: date | None
    settlement_date: date | None
    maturity_date: date | None
    isin: str | None
    principal: str | None
    currency: str | None
    security: str | None
    reopening_hint: bool
    program_code: str | None
    missing: list[str]


def _field(text: str, label: str) -> str | None:
    m = re.search(rf"^\s*{label}\s*:\s*(.+)$", text, flags=re.MULTILINE | re.IGNORECASE)
    return m.group(1).strip() if m else None


def looks_like_term_sheet(text: str) -> bool:
    return "TERM SHEET" in text.upper() and "Series" in text and "Maturity Date" in text


def parse_term_sheet(data: bytes) -> ParsedTermSheet:
    return parse_term_sheet_text(pdf_text(data))


def parse_term_sheet_text(text: str) -> ParsedTermSheet:
    series_raw = _field(text, "Series")
    series = None
    if series_raw:
        m = re.search(r"(\d{4}-\d+)", series_raw)
        series = m.group(1) if m else series_raw
    issuer = _field(text, "Issuer")
    principal = _field(text, r"(?:Aggregate\s+)?Principal Amount")
    isin = _field(text, "ISIN")
    currency = None
    if principal:
        if principal.startswith("$") or "US$" in principal or "USD" in principal:
            currency = "USD"
        else:
            m = re.match(r"([A-Z]{3}|C\$|A\$|£|€)", principal)
            currency = {"C$": "CAD", "A$": "AUD", "£": "GBP", "€": "EUR"}.get(m.group(1), m.group(1)) if m else None
    program_code = None
    if issuer:
        if "Global Funding" in issuer:
            program_code = "REPACK6883"
        elif "New York Life Funding" in issuer:
            program_code = "REPACK5970"
    reopening_hint = bool(re.search(r"(?i)re-?open|fungible|form a single series|consolidated", text))
    parsed = ParsedTermSheet(
        series=series,
        issuer=issuer,
        trade_date=parse_date(_field(text, "Trade Date")),
        settlement_date=parse_date(_field(text, r"(?:Settlement|Issue) Date")),
        maturity_date=parse_date(_field(text, r"(?:Stated\s+)?Maturity Date")),
        isin=isin,
        principal=principal,
        currency=currency,
        security=_field(text, "Security"),
        reopening_hint=reopening_hint,
        program_code=program_code,
        missing=[],
    )
    for name in ("series", "settlement_date", "maturity_date"):
        if getattr(parsed, name) is None:
            parsed.missing.append(name.replace("_", " "))
    return parsed


# --------------------------------------------------------------------------------------------
# EMTN & GMTN Funding Agreement List (the FA register)
# --------------------------------------------------------------------------------------------
def read_excel_any(data: bytes, filename: str, **kw) -> dict[str, pd.DataFrame]:
    engine = "xlrd" if filename.lower().endswith(".xls") else None
    return pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, engine=engine, **kw)


def looks_like_fa_register(sheets: dict[str, pd.DataFrame]) -> bool:
    for df in sheets.values():
        if df.shape[1] >= 5 and df.astype(str).iloc[:30, 0].str.strip().eq("FA#").any():
            return True
    return False


def parse_fa_register(data: bytes, filename: str) -> pd.DataFrame:
    """Return one clean row per FA number from the FA register workbook."""
    sheets = read_excel_any(data, filename)
    rows = []
    for df in sheets.values():
        if df.empty:
            continue
        for _, r in df.iterrows():
            fa_raw = str(r.iloc[0]).strip()
            m = re.match(r"^(\d{5,8})", fa_raw)
            if not m:
                continue  # header rows, blank rows, footers
            series_raw = str(r.iloc[1]).strip() if not pd.isna(r.iloc[1]) else ""
            notes = " ".join(str(x) for x in r.iloc[12:].tolist() if not pd.isna(x)).strip()
            reopen_of = None
            for src in (fa_raw, notes):
                mm = re.search(r"(?i)(?:original was|re-?\s?opening)\s*(\d{5})", src)
                if mm:
                    reopen_of = int(mm.group(1))
            is_reopen = bool(re.search(r"(?i)re-?\s?opening|tranche\s*\d", fa_raw + " " + series_raw + " " + notes))
            program = "REPACK5970" if "EMTN" in series_raw.upper() or "GMTN" not in series_raw.upper() else "REPACK6883"
            base = re.search(r"(\d{4}-\d+)", series_raw)
            rows.append({
                "fa_number": int(m.group(1)),
                "series_raw": series_raw,
                "base_series": base.group(1) if base else series_raw,
                "program_code": program,
                "isin": None if pd.isna(r.iloc[2]) else str(r.iloc[2]).strip(),
                "issue_date": parse_date(r.iloc[3]),
                "maturity_date": parse_date(r.iloc[4]),
                "term": None if pd.isna(r.iloc[5]) else str(r.iloc[5]).strip(),
                "principal": None if pd.isna(r.iloc[6]) else str(r.iloc[6]).strip(),
                "currency": None if pd.isna(r.iloc[7]) else str(r.iloc[7]).strip(),
                "expenses": None if pd.isna(r.iloc[10]) else str(r.iloc[10]).strip(),
                "citi_annual_expense": None if pd.isna(r.iloc[11]) else str(r.iloc[11]).strip(),
                "register_notes": notes or None,
                "is_reopening": is_reopen,
                "reopening_of": reopen_of,
            })
    out = pd.DataFrame(rows)
    if out.empty:
        raise ValueError("No FA rows found. Expected a sheet with an 'FA#' header in column A.")
    out = out.drop_duplicates("fa_number", keep="first").sort_values("fa_number").reset_index(drop=True)
    # Display series: base series plus "(n)" suffix for reopenings/tranches, matching the tracker.
    out["series"] = out["base_series"]
    for (prog, base), grp in out.groupby(["program_code", "base_series"]):
        for k, idx in enumerate(grp.sort_values("fa_number").index, start=1):
            if k > 1:
                out.loc[idx, "series"] = f"{base} ({k})"
                out.loc[idx, "is_reopening"] = True
                if pd.isna(out.loc[idx, "reopening_of"]) or out.loc[idx, "reopening_of"] is None:
                    out.loc[idx, "reopening_of"] = int(grp.sort_values("fa_number")["fa_number"].iloc[0])
    out["reopening_of"] = out["reopening_of"].astype("Int64")
    return out


# --------------------------------------------------------------------------------------------
# Legacy "Citigroup Accrued Expenses" tracker (used once, at initial load)
# --------------------------------------------------------------------------------------------
def looks_like_tracker(sheets: dict[str, pd.DataFrame]) -> bool:
    return any(re.search(r"Citigroup \d{4}", name) for name in sheets)


def parse_tracker(data: bytes, filename: str) -> pd.DataFrame:
    """Flatten the annual GMTN tabs into one row per (tab year, FA)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    rows = []
    for ws in wb.worksheets:
        m = re.search(r"(\d{4})$", ws.title.strip())
        if not m:
            continue
        year = int(m.group(1))
        for r in ws.iter_rows(min_row=1, values_only=True):
            if not r or len(r) < 8:
                continue
            series, fa = r[0], r[1]
            if not isinstance(fa, (int, float)) or series is None:
                continue
            inv = r[5]
            rows.append({
                "tab_year": year,
                "series": str(series).strip(),
                "fa_number": int(fa),
                "issue_date": parse_date(r[2]),
                "maturity_date": parse_date(r[3]),
                "amount": float(r[4]) if isinstance(r[4], (int, float)) else None,
                "invoice_no": str(int(inv)) if isinstance(inv, (int, float)) else None,
                "approved_date": parse_date(r[6]) if not isinstance(r[6], str) else None,
                "paid": bool(isinstance(inv, (int, float))),
            })
    return pd.DataFrame(rows)


def tracker_summary_total(data: bytes) -> float | None:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    if "Summary" not in wb.sheetnames:
        return None
    for row in wb["Summary"].iter_rows(values_only=True):
        if row and isinstance(row[0], str) and "total accrual" in row[0].lower():
            return float(row[1]) if isinstance(row[1], (int, float)) else None
    return None


# --------------------------------------------------------------------------------------------
# Folder / zip classification for the one-time initial load
# --------------------------------------------------------------------------------------------
@dataclass
class ClassifiedFile:
    name: str
    data: bytes
    kind: str  # invoice | term_sheet | fa_register | tracker | unknown
    detail: str = ""


def classify_file(name: str, data: bytes) -> ClassifiedFile:
    low = name.lower()
    try:
        if low.endswith(".pdf"):
            text = pdf_text(data)
            if looks_like_citi_invoice(text):
                return ClassifiedFile(name, data, "invoice")
            if looks_like_term_sheet(text):
                return ClassifiedFile(name, data, "term_sheet")
            return ClassifiedFile(name, data, "unknown", "PDF is neither a Citi invoice nor a term sheet")
        if low.endswith((".xlsx", ".xlsm", ".xls")):
            sheets = read_excel_any(data, name, nrows=40)
            if looks_like_tracker(sheets):
                return ClassifiedFile(name, data, "tracker")
            if looks_like_fa_register(sheets):
                return ClassifiedFile(name, data, "fa_register")
            return ClassifiedFile(name, data, "unknown", "Workbook layout not recognised")
    except Exception as exc:  # unreadable file
        return ClassifiedFile(name, data, "unknown", f"Could not read: {exc}")
    return ClassifiedFile(name, data, "unknown", "Unsupported file type")


def unpack_zip(data: bytes) -> list[tuple[str, bytes]]:
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir() or info.filename.startswith("__MACOSX") or info.filename.split("/")[-1].startswith("."):
                continue
            out.append((info.filename.split("/")[-1], zf.read(info)))
    return out
