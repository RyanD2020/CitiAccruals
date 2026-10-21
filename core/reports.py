"""Excel outputs. Totals on the Summary sheet are live formulas over the detail sheets."""
from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import engine
from .config import CITI_BILLING_EMAIL

DARK_BLUE, BLUE, GREY_RULE = "000A62", "0468FF", "E3E0DD"
GREEN, LIGHT_GREEN = "016355", "A5EFBF"
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=DARK_BLUE)
# Pass and Paid cells: Green on Light Green (design/BRAND.md)
PASS_FONT = Font(name="Arial", size=10, bold=True, color=GREEN)
PASS_FILL = PatternFill("solid", fgColor=LIGHT_GREEN)
TITLE_FONT = Font(name="Georgia", size=16, color=DARK_BLUE)
SUB_FONT = Font(name="Arial", size=10, color="000000")
BODY_FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
RULE = Border(bottom=Side(style="thin", color=GREY_RULE))
MONEY = '$#,##0;($#,##0);"-"'
DATE_FMT = "dd-mmm-yyyy"


def _table(ws, df: pd.DataFrame, start_row: int, money_cols=(), date_cols=(), widths=None) -> tuple[int, int]:
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, Alignment(vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            if v is not None and not isinstance(v, str):
                try:
                    if pd.isna(v):
                        v = None
                except (TypeError, ValueError):
                    pass
            if hasattr(v, "item"):
                v = v.item()
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = BODY_FONT, RULE
            if v in (engine.STATUS_PAID, "Pass"):
                c.font, c.fill = PASS_FONT, PASS_FILL
            name = df.columns[j - 1]
            if name in money_cols:
                c.number_format = MONEY
            elif name in date_cols:
                c.number_format = DATE_FMT
    for j, col in enumerate(df.columns, start=1):
        w = (widths or {}).get(col) or min(max(len(str(col)) + 2, 12), 60)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + 1, start_row + len(df)


def _detail_frame(lg: pd.DataFrame) -> pd.DataFrame:
    cols = {
        "short_name": "Program", "series": "Series", "fa_number": "FA #", "issue_date": "Issue date",
        "maturity_date": "Maturity date", "event_date": "Fee due date", "period_no": "Fee year #",
        "expected_amount": "Amount", "status": "Status", "invoice_no": "Citi invoice",
        "approval_date": "Approved", "payment_date": "Paid", "days_outstanding": "Days outstanding",
    }
    d = lg[list(cols)].rename(columns=cols).copy()
    d["FA #"] = d["FA #"].astype(int)
    d["Days outstanding"] = pd.to_numeric(d["Days outstanding"], errors="coerce")
    return d


def accrual_workbook(as_of: date, lg: engine.Ledger, roll: dict, invoices: pd.DataFrame,
                     issues: pd.DataFrame, prepared_by: str) -> bytes:
    f = lg.frame
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Citi trustee fee accrual"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"As of {as_of:%d %B %Y}. Prepared by {prepared_by}."
    ws["A2"].font = SUB_FONT

    open_df = _detail_frame(f[f["outstanding"]].sort_values(["event_date", "series"]))
    all_df = _detail_frame(f[f["event_date"] <= as_of].sort_values(["event_date", "series"]))

    ws_o = wb.create_sheet("Outstanding detail")
    first, last = _table(ws_o, open_df, 1, money_cols=("Amount",),
                         date_cols=("Issue date", "Maturity date", "Fee due date", "Approved", "Paid"),
                         widths={"Status": 30, "Series": 12})
    amt_col = get_column_letter(list(open_df.columns).index("Amount") + 1)
    st_col = get_column_letter(list(open_df.columns).index("Status") + 1)
    yr_col_idx = len(open_df.columns) + 1
    ws_o.cell(row=1, column=yr_col_idx, value="Fee year").font = HEAD_FONT
    ws_o.cell(row=1, column=yr_col_idx).fill = HEAD_FILL
    due_col = get_column_letter(list(open_df.columns).index("Fee due date") + 1)
    for r in range(first, last + 1):
        ws_o.cell(row=r, column=yr_col_idx, value=f"=YEAR({due_col}{r})").font = BODY_FONT
    yr_col = get_column_letter(yr_col_idx)
    rng = lambda col: f"'Outstanding detail'!${col}${first}:${col}${max(last, first)}"

    r = 4
    ws.cell(row=r, column=1, value="Outstanding accrual").font = BOLD
    ws.cell(row=r, column=2, value=f"=SUM({rng(amt_col)})").number_format = MONEY
    ws.cell(row=r, column=2).font = Font(name="Georgia", size=14, color=DARK_BLUE, bold=True)
    r += 2
    ws.cell(row=r, column=1, value="By status").font = Font(name="Georgia", size=12, color=BLUE)
    r += 1
    for s in engine.OPEN_STATUSES:
        ws.cell(row=r, column=1, value=s).font = BODY_FONT
        c = ws.cell(row=r, column=2, value=f'=SUMIFS({rng(amt_col)},{rng(st_col)},"{s}")')
        c.number_format, c.font = MONEY, BODY_FONT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="By fee year").font = Font(name="Georgia", size=12, color=BLUE)
    r += 1
    years = sorted(f.loc[f["outstanding"], "fee_year"].unique())
    for y in years:
        ws.cell(row=r, column=1, value=str(y)).font = BODY_FONT
        c = ws.cell(row=r, column=2, value=f"=SUMIFS({rng(amt_col)},{rng(yr_col)},{int(y)})")
        c.number_format, c.font = MONEY, BODY_FONT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Month rollforward").font = Font(name="Georgia", size=12, color=BLUE)
    r += 1
    start_r = r
    for label, key, sign in (("Opening accrual", "opening", 1), ("New fees due", "new_fees", 1),
                             ("Paid", "paid", -1), ("Written off", "written_off", -1)):
        ws.cell(row=r, column=1, value=label).font = BODY_FONT
        c = ws.cell(row=r, column=2, value=sign * roll[key])
        c.number_format, c.font = MONEY, Font(name="Arial", size=10, color="0000FF")
        r += 1
    ws.cell(row=r, column=1, value="Closing accrual").font = BOLD
    c = ws.cell(row=r, column=2, value=f"=SUM(B{start_r}:B{r - 1})")
    c.number_format, c.font = MONEY, BOLD
    r += 1
    ws.cell(row=r, column=1, value="Check against outstanding detail (should be 0)").font = SUB_FONT
    c = ws.cell(row=r, column=2, value=f"=B{r - 1}-B4")
    c.number_format = MONEY
    ws.conditional_formatting.add(f"B{r}", CellIsRule(operator="equal", formula=["0"], font=PASS_FONT, fill=PASS_FILL))
    r += 2
    ws.cell(row=r, column=1, value="Rollforward inputs (blue) come from the tool's fee ledger; totals are formulas.").font = SUB_FONT
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18

    ws_a = wb.create_sheet("All fees due")
    _table(ws_a, all_df, 1, money_cols=("Amount",),
           date_cols=("Issue date", "Maturity date", "Fee due date", "Approved", "Paid"), widths={"Status": 30})

    fu = followup_frame(lg)
    ws_f = wb.create_sheet("Citi follow-up")
    _table(ws_f, fu, 1, money_cols=("Expected fee",), date_cols=("Fee due date", "Issue date", "Maturity date"),
           widths={"Why": 50})

    ws_i = wb.create_sheet("Invoices")
    inv = invoices.sort_values("invoice_date")[["invoice_no", "program_code", "invoice_date", "arrears_start",
                                                "arrears_end", "fee_total", "line_count", "approval_date",
                                                "approved_by", "payment_date", "source"]].copy()
    inv.columns = ["Invoice", "Program", "Invoice date", "Arrears from", "Arrears to", "Total", "Lines",
                   "Approved", "Approved by", "Paid", "Source"]
    _table(ws_i, inv, 1, money_cols=("Total",),
           date_cols=("Invoice date", "Arrears from", "Arrears to", "Approved", "Paid"))

    if issues is not None and len(issues):
        ws_x = wb.create_sheet("Open exceptions")
        x = issues[issues["status"] == "Open"][["category", "entity_id", "description", "amount_impact"]].copy()
        x.columns = ["Category", "Reference", "Description", "Amount impact"]
        _table(ws_x, x, 1, money_cols=("Amount impact",), widths={"Description": 90, "Category": 28})

    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def followup_frame(lg: engine.Ledger) -> pd.DataFrame:
    f = lg.frame
    fu = f[(f["status"] == engine.STATUS_DUE)].copy()
    billed_fas = set(f.loc[f["invoice_no"].notna(), "fa_number"])
    fu["Why"] = fu.apply(
        lambda r: ("Citi has never billed this FA in the tracked period" if r["fa_number"] not in billed_fas
                   else "Not on the invoice for its quarter") + ("" if r["follow_up_due"] else " (invoice not yet due)"),
        axis=1) if len(fu) else []
    out = fu[["short_name", "series", "fa_number", "isin", "issue_date", "maturity_date", "event_date",
              "expected_amount", "days_outstanding", "Why"]].rename(columns={
        "short_name": "Program", "series": "Series", "fa_number": "FA #", "isin": "ISIN", "issue_date": "Issue date",
        "maturity_date": "Maturity date", "event_date": "Fee due date", "expected_amount": "Expected fee",
        "days_outstanding": "Days since due"})
    out["FA #"] = out["FA #"].astype(int)
    return out.sort_values(["Fee due date", "Series"]).reset_index(drop=True)


def followup_email(fu: pd.DataFrame, sender: str) -> str:
    lines = [
        f"To: {CITI_BILLING_EMAIL}",
        "Subject: New York Life Global Funding - trustee fees not yet invoiced",
        "",
        "Hello Citi GATS Billing team,",
        "",
        "Our records show the following annual trustee fees are due but have not appeared on an invoice. "
        "Could you please confirm whether they will be billed, or let us know if your records differ?",
        "",
    ]
    for r in fu.to_dict("records"):
        mat, due = r["Maturity date"], r["Fee due date"]
        lines.append(f"  - Series {r['Series']} (maturing {mat.day} {mat:%b %Y}), "
                     f"fee due {due.day} {due:%b %Y}: ${r['Expected fee']:,.0f}")
    lines += ["", f"Total: ${fu['Expected fee'].sum():,.0f}", "", "Thank you,", sender]
    return "\n".join(lines)


def xlsx_from_frames(frames: dict[str, pd.DataFrame], money_cols=(), date_cols=()) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in frames.items():
        ws = wb.create_sheet(name[:31])
        _table(ws, df, 1, money_cols=money_cols, date_cols=date_cols)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
