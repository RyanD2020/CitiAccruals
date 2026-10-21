"""The old "Citigroup Accrued Expenses" tracker, corrected row by row from the fee ledger.

The corrected workbook is a copy of the tracker file itself: same tabs, columns, formulas, colours and notes. Only
the cells that disagree with the ledger are rewritten, each one highlighted with a comment, and a "Changes from
tracker" sheet lists every row with its explanation and the business decision on it.

Tracker conventions (read from the file): column E is the year's fee (blank = not accrued yet, 0 with "N/A" in F/G
= no fee that year), F is the invoice number, G the date approved for payment, and H = ISNUMBER(F) marks it paid.
A tab's amount owed is SUM(E) minus the amounts marked paid.
"""
from __future__ import annotations

import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from . import engine
from .dates import add_years

DECISIONS = ("apply", "keep", "review")          # apply the correction / keep the tracker's value / needs more review
CHANGED_FILL = PatternFill("solid", fgColor="DCEBFF")
TOOL = "Citi fee accruals tool"

KIND_TITLE = {
    "paid_not_recorded": "Billed and paid, but not recorded on this tab",
    "wrong_tab": "Invoice recorded on the wrong year's tab",
    "not_billed": "Recorded as paid, but the invoice does not bill this FA",
    "no_fee_due": "Fee accrued, but none is due that year",
    "not_accrued": "Fee due, but not accrued",
    "not_yet_due": "Fee accrued before it is due",
    "missing_row": "FA missing from this tab",
    "other": "Other difference",
}


@dataclass
class Tab:
    title: str
    year: int
    short_name: str            # program short name from the tab title, e.g. GMTN
    first_row: int
    last_row: int
    rows: dict[int, int]       # FA number -> row
    all_row: int | None
    paid_row: int | None
    owed_row: int | None


@dataclass
class Finding:
    id: str                    # "<FA>:<year>", stable across as-of dates
    tab: str
    year: int
    row: int | None            # None for a row the tracker does not have
    fa: int
    series: str
    kind: str
    tracker: tuple             # (E, F, G) as in the tracker
    corrected: tuple           # (E, F, G) from the ledger
    effect: float              # change in the amount owed on this tab
    explanation: str
    changes: list[tuple[str, object, object]] = field(default_factory=list)

    @property
    def title(self) -> str:
        return KIND_TITLE[self.kind]


# ------------------------------------------------------------------ reading the tracker
def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _owed(e, f) -> float:
    return 0.0 if _is_num(f) or not _is_num(e) else float(e)


def _day(v):
    if isinstance(v, datetime):
        return v.date()
    return v


def read_tabs(wb) -> list[Tab]:
    tabs = []
    for ws in wb.worksheets:
        m = re.search(r"^(\w+)-.*?(\d{4})$", ws.title.strip())
        if not m:
            continue
        header = next((r for r in range(1, 40) if str(ws.cell(r, 6).value or "").strip() == "Invoice Date or #"), None)
        if header is None:
            continue
        first = header + 2
        rows, r = {}, first
        while _is_num(ws.cell(r, 2).value):
            rows[int(ws.cell(r, 2).value)] = r
            r += 1
        last = r - 1
        marks = {str(ws.cell(x, 6).value or "").strip(): x for x in range(last + 1, last + 12)}
        tabs.append(Tab(ws.title, int(ws.cell(header, 5).value or m.group(2)), m.group(1), first, last, rows,
                        marks.get("All"), marks.get("Paid"), marks.get("Owed")))
    return tabs


def tracker_total(tracker_bytes: bytes) -> float:
    wv = load_workbook(io.BytesIO(tracker_bytes), data_only=True)
    return sum(_owed(wv[t.title].cell(r, 5).value, wv[t.title].cell(r, 6).value)
               for t in read_tabs(wv) for r in range(t.first_row, t.last_row + 1))


# ------------------------------------------------------------------ what each row should say
def _ledger_frame(data: dict, as_of: date) -> pd.DataFrame:
    return engine.ledger(data["fee_events"], data["contracts"], data["invoices"], data["programs"], as_of).frame


def _desired(fa: int, year: int, lg: pd.DataFrame, invoices: pd.DataFrame, tracker_dates: dict, as_of: date):
    """(E, F, G) for one FA on one year's tab, in the tracker's conventions."""
    ev = lg[(lg["fa_number"] == fa) & (lg["fee_year"] == year)]
    if ev.empty:
        return 0, "N/A", "N/A"
    e = ev.iloc[0]
    if e["status"] == engine.STATUS_WRITTEN_OFF:
        return 0, "Written off", e["write_off_date"]
    if e["invoice_no"] and e["status"] in (engine.STATUS_PAID, engine.STATUS_APPROVED, engine.STATUS_INVOICED):
        inv = str(e["invoice_no"])
        num = int(inv) if inv.isdigit() else inv
        approved = tracker_dates.get(inv)
        if approved is None:
            row = invoices[invoices["invoice_no"] == inv]
            approved = row["approval_date"].iloc[0] if len(row) else None
        if e["status"] == engine.STATUS_PAID:
            return float(e["expected_amount"]), num, approved
        # Billed but not yet paid: the tool keeps it in the accrual, so it must not look paid (ISNUMBER = FALSE).
        return float(e["expected_amount"]), f"{inv} (not yet paid)", approved
    if e["event_date"] > as_of:
        return None, None, None          # not due yet: accrued when it falls due
    return float(e["expected_amount"]), None, None


def _same(a, b) -> bool:
    if (a in (None, "") or (isinstance(a, float) and pd.isna(a))) and (b in (None, "") or (isinstance(b, float) and pd.isna(b))):
        return True
    if _is_num(a) and _is_num(b):
        return abs(float(a) - float(b)) < 0.005
    return _day(a) == _day(b)


def _fmt(v, money: bool = False) -> str:
    v = _day(v)
    if v is None or v == "":
        return "blank"
    if isinstance(v, date):
        return f"{v.day} {v:%b %Y}"
    if _is_num(v):
        if money:
            return f"${float(v):,.0f}"
        return str(int(v)) if float(v).is_integer() else str(v)
    return str(v)


def describe(cells: tuple) -> str:
    """A tracker row's E/F/G in plain words, e.g. '$3,000, paid on invoice 167355 (15 Jan 2026)'."""
    e, f, g = cells
    if e is None and f is None:
        return "Not accrued"
    if e == 0 and f in ("N/A", None):
        return "No fee"
    amount = _fmt(e, money=True) if _is_num(e) else "No amount"
    if _is_num(f):
        return f"{amount}, paid on invoice {int(f)}" + (f" ({_fmt(g)})" if isinstance(_day(g), date) else "")
    if f:
        return f"{amount}, {f}"
    return f"{amount}, owed"


def _explain(kind, fa, year, tracker, corrected, contract, billed_by: dict, recorded_on: dict,
             first_year: int) -> str:
    te, tf, _ = tracker
    ce, cf, _ = corrected
    if kind == "paid_not_recorded":
        other = recorded_on.get((fa, cf))
        return (f"Invoice {cf} billed and paid this FA's {year} fee"
                + (f"; the tracker records it on the {other} tab instead." if other else "; the tracker does not record it."))
    if kind == "wrong_tab":
        y2, due = billed_by[(fa, str(int(tf)))]
        where = (f"so it belongs on the {y2} tab" if y2 >= first_year
                 else "which is earlier than the tracker's first tab")
        return f"Invoice {int(tf)} billed this FA's {y2} fee (due {due.day} {due:%b %Y}), {where}."
    if kind == "not_billed":
        return f"Invoice {int(tf)} does not bill this FA, so this fee is not paid."
    if kind == "no_fee_due":
        if contract is not None and contract["maturity_date"] and contract["issue_date"]:
            anniv = add_years(contract["issue_date"], year - contract["issue_date"].year)
            if contract["maturity_date"] <= anniv:
                return (f"Matures {_fmt(contract['maturity_date'])}, on or before the {_fmt(anniv)} anniversary, "
                        f"so no {year} fee is due.")
            if contract["issue_date"].year > year:
                return f"Issued {_fmt(contract['issue_date'])}, so no {year} fee is due."
        return f"No {year} fee is due under the fee schedule."
    if kind == "not_accrued":
        return f"The {year} fee is due and Citi has not billed it, so it belongs in the amount owed."
    if kind == "not_yet_due":
        return f"The {year} fee is not due yet, so it is not accrued until its due date."
    if kind == "missing_row":
        return f"The FA has a {year} fee but no row on this tab."
    return "The tracker and the fee ledger disagree on this row."


def _kind(tracker, corrected, fa: int, billed_by: dict) -> str:
    te, tf, _ = tracker
    ce, cf, _ = corrected
    if not _is_num(tf) and _is_num(cf):
        return "paid_not_recorded"
    if _is_num(tf) and not (_is_num(cf) and int(cf) == int(tf)):
        return "wrong_tab" if (fa, str(int(tf))) in billed_by else "not_billed"
    if _owed(te, tf) > 0 and ce == 0:
        return "no_fee_due"
    if _owed(te, tf) == 0 and _owed(ce, cf) > 0:
        return "not_accrued"
    if _owed(te, tf) > 0 and ce is None:
        return "not_yet_due"
    return "other"


def findings(tracker_bytes: bytes, data: dict, as_of: date) -> list[Finding]:
    """Every tracker row that disagrees with the ledger at as_of, plus FAs with a fee and no row."""
    wv = load_workbook(io.BytesIO(tracker_bytes), data_only=True)
    tabs = read_tabs(wv)
    lg = _ledger_frame(data, as_of)
    contracts = data["contracts"].set_index("fa_number")
    programs = data["programs"]
    lines = data["invoice_lines"]
    billed_by = {}  # (FA, invoice) -> (fee year, due date) for what each invoice actually billed
    for r in lines.dropna(subset=["fa_number"]).itertuples():
        billed_by[(int(r.fa_number), str(r.invoice_no))] = (r.action_date.year, r.action_date)
    dates = Counter()
    recorded_on = {}
    for t in tabs:
        ws = wv[t.title]
        for fa, r in t.rows.items():
            f, g = ws.cell(r, 6).value, _day(ws.cell(r, 7).value)
            if _is_num(f):
                recorded_on[(fa, int(f))] = t.year
                if isinstance(g, date):
                    dates[(str(int(f)), g)] += 1
    tracker_dates = {}
    for (inv, g), _ in dates.most_common():
        tracker_dates.setdefault(inv, g)

    first_year = min((t.year for t in tabs), default=0)
    out = []
    for t in tabs:
        ws = wv[t.title]
        codes = set(programs.loc[programs["short_name"] == t.short_name, "program_code"])
        for fa, r in t.rows.items():
            if fa not in contracts.index:
                continue      # not a tracked agreement (e.g. matured before tracking began)
            tracker = (ws.cell(r, 5).value, ws.cell(r, 6).value, _day(ws.cell(r, 7).value))
            corrected = _desired(fa, t.year, lg, data["invoices"], tracker_dates, as_of)
            if _same(tracker[0], corrected[0]) and _same(tracker[1], corrected[1]):
                continue
            if tracker[0] is None and tracker[1] is None and corrected[:2] == (0, "N/A"):
                continue      # left blank where no fee is due: same meaning as the tool's 0 / N/A, nothing to fix
            if _same(tracker[1], corrected[1]):
                corrected = (corrected[0], corrected[1], tracker[2])   # only the amount differs: keep their date
            kind = _kind(tracker, corrected, fa, billed_by)
            out.append(_finding(t, r, fa, contracts.loc[fa], tracker, corrected, kind, billed_by, recorded_on,
                                first_year))
        # FAs with a fee in this year and no row on the tab
        have = set(t.rows)
        due = lg[(lg["fee_year"] == t.year) & lg["program_code"].isin(codes) & ~lg["fa_number"].isin(have)]
        for fa in sorted(set(due["fa_number"].astype(int))):
            corrected = _desired(fa, t.year, lg, data["invoices"], tracker_dates, as_of)
            if corrected[0] is None:
                continue
            out.append(_finding(t, None, fa, contracts.loc[fa], (None, None, None), corrected, "missing_row",
                                billed_by, recorded_on, first_year))
    return out


def _finding(t, r, fa, c, tracker, corrected, kind, billed_by, recorded_on, first_year) -> Finding:
    names = ("Amount", "Invoice", "Date approved")
    changes = [(n, a, b) for n, a, b in zip(names, tracker, corrected) if not _same(a, b)]
    return Finding(f"{fa}:{t.year}", t.title, t.year, r, fa, str(c["series"]), kind, tracker, corrected,
                   round(_owed(*corrected[:2]) - _owed(*tracker[:2]), 2),
                   _explain(kind, fa, t.year, tracker, corrected, c, billed_by, recorded_on, first_year), changes)


# ------------------------------------------------------------------ the corrected workbook
def corrected_workbook(tracker_bytes: bytes, data: dict, as_of: date, decisions: dict[str, dict],
                       last_years: int | None = None) -> tuple[bytes, dict]:
    """Copy of the tracker with every change applied except rows decided "keep" (keep the tracker's value).

    last_years: show only the most recent N year tabs; the amount still owed on dropped tabs stays on the Summary
    as "Earlier years", because those fees remain in the accrual until they are paid or written off.
    Returns (xlsx bytes, totals).
    """
    found = findings(tracker_bytes, data, as_of)
    wb = load_workbook(io.BytesIO(tracker_bytes))
    tabs = {t.title: t for t in read_tabs(wb)}
    applied = [f for f in found if decisions.get(f.id, {}).get("decision") != "keep"]

    for f in applied:
        ws = wb[f.tab]
        t = tabs[f.tab]
        r = f.row
        if r is None:          # add the missing row at the end of the tab's FA block
            r = _append_row(ws, t, f, data["contracts"].set_index("fa_number").loc[f.fa])
        for col, (old, new) in zip((5, 6, 7), zip(f.tracker, f.corrected)):
            if _same(old, new):
                continue
            cell = ws.cell(r, col)
            cell.value = new
            cell.fill = CHANGED_FILL
            cell.comment = Comment(f"Changed from {_fmt(old, money=col == 5)}. {f.explanation}", TOOL)

    totals = _totals(wb, tabs, as_of, data)
    _summary(wb, tabs, as_of, last_years, totals)
    _changes_sheet(wb, found, decisions)
    if last_years:
        keep = sorted({t.year for t in tabs.values()}, reverse=True)[:last_years]
        for t in tabs.values():
            if t.year not in keep:
                del wb[t.title]
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), totals


def _append_row(ws, t: Tab, f: Finding, c) -> int:
    r = t.last_row + 1
    ws.insert_rows(r)
    for col in range(1, 9):     # copy the row style from the row above
        ws.cell(r, col)._style = ws.cell(r - 1, col)._style
    ws.cell(r, 1, f.series)
    ws.cell(r, 2, f.fa)
    ws.cell(r, 3, c["issue_date"])
    ws.cell(r, 4, c["maturity_date"])
    ws.cell(r, 8, f"=ISNUMBER(F{r})")
    for col in (1, 2, 3, 4):
        ws.cell(r, col).fill = CHANGED_FILL
    ws.cell(r, 1).comment = Comment(f"Row added. {f.explanation}", TOOL)
    # The total block moved down one row; rewrite its formulas for the new range.
    t.last_row += 1
    t.all_row, t.paid_row, t.owed_row = (x + 1 if x else x for x in (t.all_row, t.paid_row, t.owed_row))
    if t.all_row:
        ws.cell(t.all_row, 5, f"=SUM(E{t.first_row}:E{t.last_row})")
    if t.owed_row and t.all_row and t.paid_row:
        ws.cell(t.owed_row, 5, f"=E{t.all_row}-E{t.paid_row}")
    t.rows[f.fa] = r
    return r


def _totals(wb, tabs: dict[str, Tab], as_of: date, data: dict) -> dict:
    by_year = {}
    for t in tabs.values():
        ws = wb[t.title]
        by_year[t.year] = round(sum(_owed(ws.cell(r, 5).value, ws.cell(r, 6).value)
                                    for r in range(t.first_row, t.last_row + 1)), 2)
    lg = _ledger_frame(data, as_of)
    codes = set(data["programs"].loc[data["programs"]["short_name"].isin({t.short_name for t in tabs.values()}),
                                     "program_code"])
    ledger_owed = float(lg.loc[lg["outstanding"] & lg["program_code"].isin(codes), "expected_amount"].sum())
    return {"by_year": by_year, "corrected_total": round(sum(by_year.values()), 2), "ledger_total": ledger_owed}


def _summary(wb, tabs: dict[str, Tab], as_of: date, last_years: int | None, totals: dict) -> None:
    if "Summary" not in wb.sheetnames:
        return
    ws = wb["Summary"]
    for r in range(2, 12):
        for col in (1, 2):
            ws.cell(r, col).value = None
    ws.cell(2, 1, f"Corrected by the {TOOL} at {as_of.day} {as_of:%B %Y}. Every change is listed on the "
                  "'Changes from tracker' sheet.")
    ordered = sorted(tabs.values(), key=lambda t: t.year, reverse=True)
    shown = ordered[:last_years] if last_years else ordered
    r = 4
    for t in shown:
        ws.cell(r, 1, t.year)
        ws.cell(r, 2, f"='{t.title}'!E{t.owed_row}")
        r += 1
    earlier = round(sum(totals["by_year"][t.year] for t in ordered[len(shown):]), 2)
    if earlier:
        ws.cell(r, 1, "Earlier years")
        ws.cell(r, 2, earlier)
        ws.cell(r, 3, "Still in the accrual until paid or written off; their tabs are not shown.")
        r += 1
    ws.cell(r, 1, "Total accrual")
    ws.cell(r, 2, f"=SUM(B4:B{r - 1})")


def _changes_sheet(wb, found: list[Finding], decisions: dict[str, dict]) -> None:
    ws = wb.create_sheet("Changes from tracker", 1)
    head = ["Tab", "Row", "Series #", "FA #", "Change", "Field", "Tracker value", "Corrected value",
            "Effect on amount owed", "Why", "Decision", "Decided by", "Note"]
    ws.append(head)
    label = {"apply": "Apply", "keep": "Keep current (not applied)", "review": "Needs more review (applied)"}
    for f in sorted(found, key=lambda x: (-x.year, x.kind, x.fa)):
        d = decisions.get(f.id, {})
        for i, (fld, old, new) in enumerate(f.changes or [("Row", None, "added")]):
            ws.append([f.year, f.row or "added", f.series, f.fa, f.title, fld, _fmt(old, fld == "Amount"),
                       _fmt(new, fld == "Amount"),
                       f.effect if i == 0 else None, f.explanation if i == 0 else None,
                       label.get(d.get("decision"), "Suggested, not yet decided (applied)"),
                       d.get("by"), d.get("note")])
    for col, w in zip("ABCDEFGHIJKLM", (8, 8, 12, 9, 40, 14, 16, 18, 12, 70, 30, 28, 40)):
        ws.column_dimensions[col].width = w
    for (cell,) in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        cell.number_format = '$#,##0;-$#,##0;"-"'
    ws.freeze_panes = "A2"
